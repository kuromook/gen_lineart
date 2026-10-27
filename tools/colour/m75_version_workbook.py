#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""**追補103 —— 版の判断を作者に仰ぐ Excel を組む(作者の裁定 2026-10-27 を実装する)。**

**作者の裁定(そのまま写す)**:

1. **ページでハッシュ値が同じなら同じファイルという判断でよい**
2. **どちらが古いか新しいかは、フォルダ名から判断できる年代でまず判断してよい**
   (**年代が進んだフォルダのものの方が古いことはないとみていい。あったとしたら
   事故のレベルだが、一旦はその判断とする**)
3. **同年代フォルダ下で同じページとみなせるものは、作者に問い合わせるしかない。**
   **UI ではやりにくい。Excel にレイヤ数などの構成を列に入れて判断を仰ぐ形にする**
   (**理由: 作者の側でも見落とす可能性があるので、ファイル情報や統計量で見逃しを抑えたい**)

**こちらは決めない。** 器具は**分けて、並べて、違う列に印を付ける**だけである(規律6)。

出力は **`data/` 配下**(原稿由来。`.gitignore` 済み。**絶対にコミットしない**)。

  <venv>/python tools/colour/m75_version_workbook.py
"""
import collections
import csv
import datetime
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

H = pathlib.Path('../lineart/inbox')
DUP = H / 'sha256_dup_clusters_20261026.csv'          # 署名の116塊 / 248頁
ALLH = H / 'request9_file_sha256_all_20261026.csv'    # 全数 9,653(33件を含む)
STRU = H / 'census_r8_layer_structure_20261027.csv'   # 層構造 437,561行
OUT = pathlib.Path('data/version_workbook')

# **種類の列**(`v_layer_kind` の値。器具が作品名を持たないようにする)
KINDS = ['raster', 'mask_fill', 'folder', 'text', 'paper', 'frame_folder',
         'story_info', 'image_material', 'gradation', 'filter', 'vector',
         '3d', 'other']
KIND_JA = {'raster': 'ラスタ', 'mask_fill': 'mask_fill', 'folder': 'フォルダー',
           'text': 'テキスト', 'paper': '用紙', 'frame_folder': 'コマ枠',
           'story_info': '作品情報', 'image_material': '画像素材',
           'gradation': 'グラデ', 'filter': '補正層', 'vector': 'ベクター',
           '3d': '3D', 'other': 'その他'}
YEAR = re.compile(r'^(\d{4})([a-z]?)$')


def rel(p):
    """根の道を落とす(**器組に根を書かない**)。"""
    low = p.lower()
    i = low.find('manuscripts\\')
    return p[i + len('manuscripts\\'):] if i >= 0 else p


def year_of(p):
    """**年代 = `manuscripts\\` の次の成分**。`2022b` は `2022` と同じ年代に置く。"""
    head = rel(p).split('\\')[0]
    m = YEAR.match(head)
    return (int(m.group(1)), m.group(2)) if m else None


def load_structure(want):
    """頁ごとの構成(**要る頁だけ**。51MB を全部持たない)。"""
    agg = collections.defaultdict(collections.Counter)
    for r in csv.DictReader(STRU.open(encoding='utf-8')):
        p = r['page_path']
        if p not in want:
            continue
        a = agg[p]
        a['層'] += 1
        a['非フォルダー'] += 0 if r['is_folder'] in ('True', '1', 'true') else 1
        a[r['v_layer_kind']] += 1
        if not (int(r['visibility'] or 0) & 1):
            a['非表示'] += 1
        if int(r['clip'] or 0) == 1:
            a['クリップ'] += 1
        if int(r['opacity'] or 256) < 256:
            a['不透明度<256'] += 1
        a['深さ'] = max(a['深さ'], int(r['depth'] or 0))
    return agg


COLS = (['相対パス', '年代', '塊', '役割', '**同じ中身の相手(裁定1)**',
         'バイト数', '更新時刻', 'sha256(先頭12)',
         '層', '非フォルダー', '非表示', 'クリップ', '不透明度<256', '深さ']
        + [KIND_JA[k] for k in KINDS]
        + ['**この塊の中で違う列**', '【判断】使う方に ○', '【判断】理由・備考'])
CMP = (['バイト数', '層', '非フォルダー', '非表示', 'クリップ', '不透明度<256', '深さ']
       + [KIND_JA[k] for k in KINDS])


def row_of(p, d, s, cid):
    out = {'相対パス': rel(p), '年代': (lambda y: f'{y[0]}{y[1]}' if y else '年不明')(year_of(p)),
           '塊': cid, '役割': d.get('version_role', ''),
           'バイト数': int(d['file_bytes'] or 0), '更新時刻': d.get('f_mtime', ''),
           'sha256(先頭12)': (d.get('sha256') or '')[:12],
           '層': s['層'], '非フォルダー': s['非フォルダー'], '非表示': s['非表示'],
           'クリップ': s['クリップ'], '不透明度<256': s['不透明度<256'],
           '深さ': s['深さ']}
    out.update({KIND_JA[k]: s[k] for k in KINDS})
    return out


def write_sheet(ws, rows, note, fill_diff=True):
    yellow = PatternFill('solid', fgColor='FFF2A8')
    grey = PatternFill('solid', fgColor='EDEDED')
    ws.append([note])
    ws.cell(1, 1).font = Font(bold=True)
    ws.append(COLS)
    for c in range(1, len(COLS) + 1):
        ws.cell(2, c).font = Font(bold=True)
        ws.cell(2, c).alignment = Alignment(wrap_text=True, vertical='top')
    by_cl = collections.defaultdict(list)
    for r in rows:
        by_cl[r['塊']].append(r)
    at = 3
    for cid in sorted(by_cl, key=lambda c: (-len(by_cl[c]), str(c))):
        mem = by_cl[cid]
        # **裁定1 をこの表の中でも効かせる**: 同じ sha256 の相手を名指しする。
        # **同じ中身なら、作者が選び直す必要は無い**
        for m in mem:
            twin = [o['相対パス'] for o in mem
                    if o is not m and o['sha256(先頭12)'] == m['sha256(先頭12)']]
            m['**同じ中身の相手(裁定1)**'] = (' ;; '.join(twin) if twin
                                            else '(この塊に同じ中身は無い)')
        diff = [c for c in CMP if len({m[c] for m in mem}) > 1]
        for m in mem:
            m['**この塊の中で違う列**'] = ' / '.join(diff) if diff else '(全部同じ)'
            ws.append([m.get(c, '') for c in COLS])
            if fill_diff:
                for c in diff:
                    ws.cell(at, COLS.index(c) + 1).fill = yellow
            at += 1
        ws.append([''] * len(COLS))          # 塊の区切り
        for c in range(1, len(COLS) + 1):
            ws.cell(at, c).fill = grey
        at += 1
    for i, c in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (
            46 if c == '相対パス' else 46 if c.startswith('**同じ中身')
            else 30 if c.startswith('**') or c.startswith('【')
            else 13)
    ws.freeze_panes = 'A3'


def main():
    dup = list(csv.DictReader(DUP.open(encoding='utf-8')))
    allh = {r['page_path']: r for r in csv.DictReader(ALLH.open(encoding='utf-8'))}
    print(f'材料  重複の表 {len(dup)}頁 / {len({r["dup_cluster_id"] for r in dup})}塊'
          f'  全数ハッシュ {len(allh)}')

    want = {r['page_path'] for r in dup}
    stru = load_structure(want)
    print(f'層構造が引けた頁 {len(stru)} / {len(want)}')

    for r in dup:                             # 全数の側の値を足す(mtime)
        a = allh.get(r['page_path'])
        if a:
            r['f_mtime'] = a['f_mtime']
            if not r.get('file_bytes'):
                r['file_bytes'] = a['f_bytes']

    by_cl = collections.defaultdict(list)
    for r in dup:
        by_cl[r['dup_cluster_id']].append(r)

    ask, byyear, same, uncov = [], [], [], []
    tally = collections.Counter()
    for cid, mem in by_cl.items():
        shas = {r['sha256'] for r in mem}
        years = {year_of(r['page_path']) for r in mem}
        rows = [row_of(r['page_path'], r, stru[r['page_path']], cid)
                for r in mem if r['page_path'] in stru]
        if len(rows) != len(mem):
            uncov += [{'相対パス': rel(r['page_path']), '塊': cid,
                       '理由': '層構造が引けない'} for r in mem
                      if r['page_path'] not in stru]
        if len(mem) < 2:
            tally['頁が1つ(塊でない)'] += 1
            continue
        if len(shas) == 1:
            tally['**裁定1: バイト同一 = 同じファイル**'] += 1
            same += rows
        elif None in years:
            tally['年代が読めない -> 問い'] += 1
            ask += rows
        elif len({y[0] for y in years}) > 1:
            tally['**裁定2: 年代で決まる**'] += 1
            byyear += rows
        else:
            tally['**裁定3: 同年代 -> 作者に問う**'] += 1
            ask += rows

    print('\n== 塊の行き先(理由別・残り 0)==')
    n = sum(tally.values())
    for k, v in tally.most_common():
        print(f'  {k:34s} {v:>4}')
    print(f'  {"合計":34s} {n:>4}   **残り {len(by_cl) - n}**')

    # **更新時刻は一括コピーの産物**(全 9,653 の 65.6% が同じ日付)。
    # 裁定2 を覆す証拠には使えない。**使えない理由も数で書く**
    bulkday = collections.Counter(
        r['f_mtime'][:10] for r in allh.values() if r['f_mtime']).most_common(1)[0]
    print(f'\n== 更新時刻の素性 ==\n  最も多い日付 {bulkday[0]} が {bulkday[1]} 件'
          f' = {bulkday[1] / len(allh):.1%} —— **一括コピーの産物。年代を覆す証拠にはしない**')

    pairs, undec = [], 0
    for cid, mem in by_cl.items():
        ok = [(year_of(r['page_path']), r.get('f_mtime', ''), rel(r['page_path']))
              for r in mem if year_of(r['page_path'])]
        for i in range(len(ok)):
            for j in range(len(ok)):
                (ya, ma, pa), (yb, mb, pb) = ok[i], ok[j]
                if ya[0] >= yb[0] or not (ma and mb):
                    continue
                if ma[:10] == bulkday[0] or mb[:10] == bulkday[0]:
                    undec += 1
                    continue
                if ma > mb:            # 古い年代の方が、更新時刻が新しい
                    pairs.append({'塊': cid, '古い年代の道': pa, '年代a': ya[0],
                                  '更新時刻a': ma, '新しい年代の道': pb,
                                  '年代b': yb[0], '更新時刻b': mb})
    print(f'  年代と更新時刻が食い違う対 **{len(pairs)}**'
          f' / 一括コピーの日付が混ざって判定できない対 {undec}')

    # 全数の中のバイト同一(**33件を含む。署名の塊に依らない**)
    byh = collections.defaultdict(list)
    for p, r in allh.items():
        if r['f_sha256']:
            byh[r['f_sha256']].append(p)
    groups = {h: ps for h, ps in byh.items() if len(ps) > 1}
    in_dup = {r['page_path'] for r in dup}
    outside = [(h, ps) for h, ps in groups.items()
               if not any(p in in_dup for p in ps)]
    print(f'\n== 全数 {len(allh)} のバイト同一の組 {len(groups)} '
          f'(頁 {sum(len(v) for v in groups.values())})==')
    print(f'  **署名の116塊の外に在る組 {len(outside)}** '
          f'(頁 {sum(len(v) for _, v in outside)}) —— **33件を含む表で数えた**')

    OUT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.datetime.now().strftime('%Y%m%d')
    path = OUT / f'version_decide_{stamp}.xlsx'
    wb = Workbook()

    ws = wb.active
    ws.title = '読み方'
    for line in [
        ['版の判断(作者に仰ぐ表)', ''],
        ['作った日時',
         datetime.datetime.now().strftime('%Y-%m-%d %H:%M')],
        ['', ''],
        ['作者の裁定(2026-10-27)', ''],
        ['1', 'ページでハッシュ値が同じなら同じファイル'],
        ['2', '古い/新しいは、フォルダ名の年代でまず判断してよい'
              '(年代が進んだ方が古いことはない。あれば事故)'],
        ['3', '同年代フォルダ下で同じページとみなせるものは作者に問う'],
        ['', ''],
        ['この表がしていること', '分けて、並べて、塊の中で違う列に色を付けた。'
                                '**どちらを使うかは決めていない**'],
        ['記入する列', '「【判断】使う方に ○」と「【判断】理由・備考」の2列'],
        ['色', '黄色 = その塊の中で値が違う列(見落としを抑えるため)'],
        ['', ''],
        ['シート', ''],
        ['問い(同年代)', f'**ここを見てください。** {len(ask)} 行'],
        ['年で決めた', f'裁定2 で決まった分。確認だけ。{len(byyear)} 行'],
        ['年代と時刻の食い違い',
         f'**{len(pairs)} 対**。更新時刻の 65.6% は一括コピー(2022-05-06)の'
         f'産物なので、**年代を覆す証拠にはしていない**。'
         f'判定できない対 {undec}'],
        ['同一ファイル', f'裁定1 で消えた分。{len(same)} 行'],
        ['バイト同一(全数)', f'9,653ファイル全部で数えたバイト同一の組 '
                             f'{len(groups)}(署名の塊の外に {len(outside)})'],
        ['覆っていない', f'{len(uncov)} 件'],
        ['', ''],
        ['材料', '重複の表 248頁/116塊(署名)・全数ハッシュ 9,653・層構造 437,561行'],
        ['注意', '重複の表は在庫 7,634頁から組んであり、33件(ol 29 + 年直下 4)が'
                 '外に在った。「バイト同一(全数)」はその33件を含む表で数えている'],
    ]:
        ws.append(line)
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 96
    for r in range(1, ws.max_row + 1):
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(1, 1).font = Font(bold=True, size=14)

    write_sheet(wb.create_sheet('問い(同年代)'), ask,
                '**同年代フォルダ下。どちらを使うか、塊ごとに ○ を付けてください**'
                '(黄色 = その塊の中で違う列)')
    write_sheet(wb.create_sheet('年で決めた'), byyear,
                '裁定2(年代が進んだ方が新しい)で決まった分。**確認だけ**')
    print(f'\n== Excel の行数 ==\n  問い(同年代) {len(ask)} 行 /'
          f' 年で決めた {len(byyear)} 行 / 同一ファイル {len(same)} 行')
    write_sheet(wb.create_sheet('同一ファイル'), same,
                '裁定1(バイト同一)で消えた分。**問いではない**')

    ws = wb.create_sheet('年代と時刻の食い違い')
    ws.append([f'**古い年代の方が、更新時刻が新しい対**({len(pairs)} 対)'
               ' —— 作者の「あったとしたら事故のレベル」に当たる形'])
    ws.append([f'ただし**更新時刻の {bulkday[1] / len(allh):.0%} は一括コピー'
               f'({bulkday[0]})の産物**。片方でもその日付の対は'
               f'「判定できない」として外した({undec} 対)。'
               '**更新時刻は年代を覆す証拠にしない**'])
    ws.append(['塊', '古い年代の道', '年代a', '更新時刻a',
               '新しい年代の道', '年代b', '更新時刻b'])
    for a in pairs:
        ws.append([a['塊'], a['古い年代の道'], a['年代a'], a['更新時刻a'],
                   a['新しい年代の道'], a['年代b'], a['更新時刻b']])
    if not pairs:
        ws.append(['(無し)'] + [''] * 6)
    ws.column_dimensions['B'].width = 52
    ws.column_dimensions['E'].width = 52

    ws = wb.create_sheet('バイト同一(全数)')
    ws.append(['**全数 9,653ファイルで数えたバイト同一の組**'
               '(33件を含む。署名の塊に依らない)'])
    ws.append(['sha256(先頭12)', '頁数', '署名の塊の中にあるか', '相対パス(;; 区切り)'])
    for h, ps in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        ws.append([h[:12], len(ps),
                   'ある' if any(p in in_dup for p in ps) else '**ない**',
                   ' ;; '.join(rel(p) for p in sorted(ps))])
    ws.column_dimensions['D'].width = 110

    ws = wb.create_sheet('覆っていない')
    ws.append(['理由別。**残り 0 にする**'])
    ws.append(['相対パス', '塊', '理由'])
    for u in uncov:
        ws.append([u['相対パス'], u['塊'], u['理由']])
    if not uncov:
        ws.append(['(無し)', '', ''])
    ws.column_dimensions['A'].width = 52

    wb.save(path)
    print(f'\nwrote {path}')
    print('  **`data/` 配下。`.gitignore` 済み。コミットしない**')
    return 0


if __name__ == '__main__':
    sys.exit(main())
